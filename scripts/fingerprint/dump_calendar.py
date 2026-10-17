"""Validate official 2026 calendars and emit deterministic fingerprint input."""
from collections import Counter
import datetime
import sys
import re
import zipfile
from xml.etree.ElementTree import ParseError
import openpyxl
from openpyxl.utils.exceptions import InvalidFileException


class SourceUnavailableError(Exception):
    pass


def dump_calendar(path):
    if not zipfile.is_zipfile(path):
        raise SourceUnavailableError('Not a valid XLSX/ZIP file')
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, ParseError, KeyError, ValueError, EOFError) as error:
        raise SourceUnavailableError(f'Invalid XLSX source: {error}') from error
    if not wb.sheetnames:
        raise SourceUnavailableError('No worksheets')
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    expected_headers = ('Runde', 'Dato', 'Dag', 'Tid', 'Hjemmelag', 'Resultat', 'Bortelag', 'Bane', 'Turnering', 'Kampnummer')
    if not rows or tuple(rows[0][:10]) != expected_headers:
        raise SourceUnavailableError('Invalid calendar headers')
    matches = [r for r in rows[1:] if any(c is not None and c != '' for c in r)]
    ids, teams = set(), set()
    for row in matches:
        if len(row) < 10 or not isinstance(row[1], (datetime.date, datetime.datetime)) or not all(row[i] for i in (4, 6, 9)):
            raise SourceUnavailableError('Incomplete calendar row')
        if not re.fullmatch(r'\d+', str(row[0])) or not re.fullmatch(r'([01]\d|2[0-3]):[0-5]\d', str(row[3])) or not re.fullmatch(r'\d+', str(row[9]).strip()):
            raise SourceUnavailableError('Invalid round, time or match ID')
        if row[5] and not re.fullmatch(r'\d+\s*-\s*\d+', str(row[5]).strip()):
            raise SourceUnavailableError('Invalid match result')
        mid = str(row[9]).strip()
        if mid in ids:
            raise SourceUnavailableError('Duplicate calendar match')
        ids.add(mid)
        teams.update(str(row[i]).strip() for i in (4, 6))
    expected = 12 if str(path).endswith('205403.xlsx') else 10
    if len(teams) != expected or len(matches) != expected * (expected - 1):
        raise SourceUnavailableError('Incomplete tournament calendar')
    pairs = Counter(tuple(sorted((str(r[4]).strip(), str(r[6]).strip()))) for r in matches)
    if len(pairs) != expected * (expected - 1) // 2 or any(a == b or count != 2 for (a, b), count in pairs.items()):
        raise SourceUnavailableError('Incomplete or duplicate tournament matchups')
    return '\n'.join(str(path) + ' ' + '|'.join('' if c is None else str(c) for c in row) for row in rows)


def main():
    print('\n'.join(dump_calendar(path) for path in sys.argv[1:]))


if __name__ == '__main__':
    main()
